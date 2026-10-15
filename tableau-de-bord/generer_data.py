#!/usr/bin/env python3
"""Génère tableau-de-bord/data.json à partir d'un export .xlsx de la base Civipédia.

Usage :
    python3 generer_data.py BASE.xlsx [--sortie data.json] [--date AAAA-MM-JJ] [--exclure OPP-0005 ...]

BASE.xlsx : export complet du classeur JOURNAL_TRAITEMENT_CIVIPEDIA_LAB
(Google Sheets → Fichier → Télécharger → Microsoft Excel). Ce fichier contient
des données privées : ne jamais le déposer dans ce dépôt.

Principe : la base est la source de vérité ; data.json est une projection
publique (comptes agrégés, événements à fiche vérifiée, appels qui reviennent).
Le script lit les colonnes par leur nom (jamais par leur position), n'invente
rien et signale les données qu'il écarte. Il ne lit que la base Civipédia :
aucune donnée de la base Crossac (POC) n'entre dans data.json.

Dépendance : openpyxl (pip install openpyxl).
"""
import argparse
import datetime as dt
import json
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl manquant : pip install openpyxl")

PREFIXES = ("MASE-C", "MASE-O", "MASE-T", "MASE-E", "MASE-P", "MASE-V", "OPP")
TABLES = {  # onglet → colonne identifiant
    "Contacts": "ID_contact",
    "Organisations": "ID_organisation",
    "Territoires": "ID_territoire",
    "Expertises": "ID_expertise",
    "Projets": "ID_projet",
    "Événements": "ID_evenement",
    "Opportunites": "ID_OPP",
    "Relations": "ID_relation",
    "Sources": "ID_source",
}
ID_INTERNE = re.compile(r"\b(MASE-[A-Z]-\d{3,4}|OPP-\d{4}|APP-\d{3}|CR-[A-Z]+-\d{3}|VEI-\d{8}-\d{3})\b")
CLES = {
    "racine": ["instantane", "genere_le", "regles", "chiffres", "evenements", "appels", "territoires", "expertises"],
    "chiffres": ["objet", "nombre"],
    "evenements": ["titre", "debut", "fin", "lieu", "territoire", "organisateur", "src"],
    "appels": ["titre", "organisme", "objet", "derniere_edition", "montant", "src"],
    "territoires": ["nom", "type", "organisations", "personnes", "projets", "evenements"],
    "expertises": ["libelle", "organisations", "personnes"],
}

alertes = []


def alerte(msg):
    alertes.append(msg)


def txt(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def date_iso(v):
    """JJ/MM/AAAA ou date Excel → AAAA-MM-JJ ; None si le jour est inconnu ou illisible."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", txt(v))
    if not m:
        return None
    try:
        return dt.date(int(m[3]), int(m[2]), int(m[1])).isoformat()
    except ValueError:
        return None


def ids(v):
    return [x.strip() for x in txt(v).split(";") if x.strip()]


def prefixe(i):
    return next((p for p in PREFIXES if i.startswith(p + "-")), None)


def majuscule(s):
    return s[:1].upper() + s[1:] if s else s


def lire(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
    base = {}
    for onglet, cle in TABLES.items():
        if onglet not in wb.sheetnames:
            sys.exit(f"Onglet absent de l'export : {onglet}")
        lignes = list(wb[onglet].iter_rows(values_only=True))
        entetes = [txt(h) for h in lignes[0]]
        if cle not in entetes:
            sys.exit(f"Colonne {cle} absente de l'onglet {onglet}")
        base[onglet] = [dict(zip(entetes, l)) for l in lignes[1:] if txt(l[entetes.index(cle)])]
    return base


def source_publique(champ, sources, contexte):
    for i in ids(champ):
        s = sources.get(i)
        if s and txt(s.get("URL_source")):
            return {"nom": txt(s.get("Nom_source")), "url": txt(s.get("URL_source"))}
    alerte(f"{contexte} : aucune source avec URL ({txt(champ) or 'Source vide'})")
    return None


def construire(base, jour, exclure):
    idx = {o: {txt(r[TABLES[o]]): r for r in base[o]} for o in TABLES}
    existants = set().union(*(idx[o].keys() for o in TABLES if o not in ("Relations", "Sources")))
    sources = idx["Sources"]

    # Relations valides : les deux ID existent, type et source renseignés.
    valides, hypotheses = [], 0
    for r in base["Relations"]:
        rid, de, vers = txt(r["ID_relation"]), txt(r.get("ID_de")), txt(r.get("ID_vers"))
        if de not in existants or vers not in existants:
            alerte(f"{rid} écartée : identifiant introuvable ({de} → {vers})")
            continue
        if not txt(r.get("Type_relation")) or not txt(r.get("Source")):
            alerte(f"{rid} écartée : Type_relation ou Source vide")
            continue
        if txt(r.get("Nature_info")).upper().startswith("HYPOTH"):
            hypotheses += 1
            continue
        valides.append(r)

    def relies(type_a, type_b):
        """Paires (objet de type_a, objet de type_b) reliées, dans un sens ou l'autre.
        Les liens « à valider » (Notes) sont exclus, comme les HYPOTHÈSE."""
        paires = set()
        for r in valides:
            if "à valider" in txt(r.get("Notes")).lower():
                continue
            de, vers = txt(r["ID_de"]), txt(r["ID_vers"])
            if prefixe(de) == type_a and prefixe(vers) == type_b:
                paires.add((de, vers))
            elif prefixe(vers) == type_a and prefixe(de) == type_b:
                paires.add((vers, de))
        return paires

    def compte(paires, cible):
        return len({a for a, b in paires if b == cible})

    # Chiffres clés
    chiffres = [
        {"objet": "Personnes", "nombre": len(base["Contacts"])},
        {"objet": "Organisations", "nombre": len(base["Organisations"])},
        {"objet": "Territoires", "nombre": len(base["Territoires"])},
        {"objet": "Projets", "nombre": len(base["Projets"])},
        {"objet": "Événements", "nombre": len(base["Événements"])},
        {"objet": "Relations documentées", "nombre": len(valides)},
    ]

    # Événements publiés : fiche « Vérifié », dates au jour près, non annulés.
    evenements = []
    org_evt = {}
    for r in valides:
        t = txt(r["Type_relation"]).lower()
        if t in ("organisateur", "co-organisateur") and prefixe(txt(r["ID_de"])) == "MASE-O" and prefixe(txt(r["ID_vers"])) == "MASE-V":
            org_evt.setdefault(txt(r["ID_vers"]), []).append((t != "organisateur", txt(idx["Organisations"][txt(r["ID_de"])].get("Nom"))))
    for e in base["Événements"]:
        eid = txt(e["ID_evenement"])
        if txt(e.get("Statut")) != "Vérifié":
            continue
        if txt(e.get("Statut_evenement")).lower().startswith("annul"):
            continue
        debut, fin = date_iso(e.get("Date_debut")), date_iso(e.get("Date_fin")) or date_iso(e.get("Date_debut"))
        if not debut:
            alerte(f"{eid} vérifié mais non publié : Date_debut illisible ({txt(e.get('Date_debut'))})")
            continue
        terr = idx["Territoires"].get(txt(e.get("ID_territoire")))
        evenements.append({
            "titre": txt(e.get("Nom")),
            "debut": debut,
            "fin": fin,
            "lieu": txt(e.get("Lieu_source")),
            "territoire": txt(terr.get("Nom_territoire")) if terr else "",
            "organisateur": " ; ".join(n for _, n in sorted(org_evt.get(eid, []))),
            "src": source_publique(e.get("Source"), sources, eid),
        })
    evenements.sort(key=lambda x: (x["debut"], x["titre"]))

    # Appels qui reviennent : Etat_appel « Clos — à guetter », information FAIT ou VÉRIFIÉ.
    appels = []
    for o in base["Opportunites"]:
        oid = txt(o["ID_OPP"])
        etat = txt(o.get("Etat_appel")).lower()
        if not (etat.startswith("clos") and "guetter" in etat):
            continue
        if oid in exclure:
            alerte(f"{oid} exclu à la demande (--exclure)")
            continue
        if txt(o.get("Nature_info")).upper() not in ("FAIT", "VÉRIFIÉ", "VERIFIE"):
            alerte(f"{oid} non publié : Nature_info = {txt(o.get('Nature_info')) or 'vide'}")
            continue
        src = source_publique(o.get("Source"), sources, oid)
        if not src:
            continue
        appels.append({
            "titre": txt(o.get("Titre")),
            "organisme": txt(o.get("Organisme_emetteur")),
            "objet": txt(o.get("Objet_conditions")),
            "derniere_edition": majuscule(txt(o.get("Derniere_edition"))),
            "montant": txt(o.get("Montant")),
            "src": src,
        })

    # Territoires : organisations et personnes reliées (Relations), projets et
    # événements dont c'est le territoire principal (ID_territoire) ou reliés.
    ot, ct, pt, vt = (relies(p, "MASE-T") for p in ("MASE-O", "MASE-C", "MASE-P", "MASE-V"))
    for o, k in (("Projets", "ID_projet"), ("Événements", "ID_evenement")):
        cible = pt if o == "Projets" else vt
        for r in base[o]:
            t = txt(r.get("ID_territoire"))
            if t:
                if t in idx["Territoires"]:
                    cible.add((txt(r[k]), t))
                else:
                    alerte(f"{txt(r[k])} : ID_territoire inconnu ({t})")
    territoires = [{
        "nom": txt(t.get("Nom_territoire")),
        "type": txt(t.get("Type_territoire")),
        "organisations": compte(ot, tid),
        "personnes": compte(ct, tid),
        "projets": compte(pt, tid),
        "evenements": compte(vt, tid),
    } for tid, t in idx["Territoires"].items()]

    oe, ce = relies("MASE-O", "MASE-E"), relies("MASE-C", "MASE-E")
    expertises = [{
        "libelle": txt(e.get("Libellé")),
        "organisations": compte(oe, eid),
        "personnes": compte(ce, eid),
    } for eid, e in idx["Expertises"].items()]

    data = {
        "instantane": f"Base Civipédia, export du {jour.strftime('%d/%m/%Y')}",
        "genere_le": jour.isoformat(),
        "regles": ("Projection publique : seuls les événements dont la fiche est vérifiée, des appels décrits "
                   "par leur source officielle, et des comptes agrégés. Aucun nom de personne, aucune "
                   "coordonnée, aucun identifiant interne."),
        "chiffres": chiffres,
        "evenements": evenements,
        "appels": appels,
        "territoires": territoires,
        "expertises": expertises,
    }
    stats = {"relations_lignes": len(base["Relations"]), "relations_valides": len(valides), "relations_hypotheses": hypotheses}
    return data, stats


def controler(data, base):
    """Contrat attendu par tableau-de-bord/index.html + absence de données privées."""
    erreurs = []
    if list(data) != CLES["racine"]:
        erreurs.append("clés racine différentes du contrat")
    for liste in ("chiffres", "evenements", "appels", "territoires", "expertises"):
        for i, x in enumerate(data[liste]):
            if list(x) != CLES[liste]:
                erreurs.append(f"{liste}[{i}] : clés {list(x)}")
    for e in data["evenements"]:
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", e["debut"]) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", e["fin"]) or e["fin"] < e["debut"]:
            erreurs.append(f"dates invalides : {e['titre']}")
    for liste, champs in (("territoires", ("organisations", "personnes", "projets", "evenements")), ("expertises", ("organisations", "personnes"))):
        for x in data[liste]:
            if any(not isinstance(x[c], int) for c in champs):
                erreurs.append(f"{liste} : compte non entier")
    texte = json.dumps(data, ensure_ascii=False)
    for m in sorted(set(ID_INTERNE.findall(texte))):
        erreurs.append(f"identifiant interne publié : {m}")
    for c in base["Contacts"]:
        nom, prenom = txt(c.get("Nom")), txt(c.get("Prénom"))
        if nom and prenom and (f"{prenom} {nom}" in texte or f"{nom} {prenom}" in texte):
            erreurs.append(f"nom de personne publié : {txt(c.get('ID_contact'))}")
        for coord in re.split(r"\s*;\s*", txt(c.get("Coordonnées"))):
            if len(coord) > 5 and coord in texte:
                erreurs.append(f"coordonnée publiée : {txt(c.get('ID_contact'))}")
    return erreurs


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("export_xlsx")
    ap.add_argument("--sortie", default="data.json")
    ap.add_argument("--date", help="date de l'export (AAAA-MM-JJ), par défaut aujourd'hui")
    ap.add_argument("--exclure", nargs="*", default=[], help="ID d'opportunités à ne pas publier")
    a = ap.parse_args()
    jour = dt.date.fromisoformat(a.date) if a.date else dt.date.today()
    base = lire(a.export_xlsx)
    data, stats = construire(base, jour, set(a.exclure))
    erreurs = controler(data, base)
    print(f"Relations : {stats['relations_lignes']} lignes, {stats['relations_valides']} valides publiées en compte, "
          f"{stats['relations_hypotheses']} HYPOTHÈSE non comptées", file=sys.stderr)
    print(f"Publié : {len(data['evenements'])} événements, {len(data['appels'])} appels, "
          f"{len(data['territoires'])} territoires, {len(data['expertises'])} expertises", file=sys.stderr)
    for m in alertes:
        print("ALERTE  " + m, file=sys.stderr)
    if erreurs:
        for m in erreurs:
            print("ERREUR  " + m, file=sys.stderr)
        sys.exit("data.json NON écrit : contrat ou confidentialité non respectés.")
    with open(a.sortie, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print(f"Écrit : {a.sortie}", file=sys.stderr)


if __name__ == "__main__":
    main()
